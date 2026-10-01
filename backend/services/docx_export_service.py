from io import BytesIO

from docx import Document

from schemas.research_report import (
    ResearchReportResponse,
)


def generate_research_report_docx(
    report: ResearchReportResponse,
) -> bytes:
    """
    Generates a Microsoft Word (.docx) document
    from an existing research report.

    This service does NOT build report data.
    It only formats the existing report.
    """

    document = Document()

    # -------------------------------------------------
    # Title
    # -------------------------------------------------

    document.add_heading(
        "AI Research Report",
        level=1,
    )

    # -------------------------------------------------
    # Project Information
    # -------------------------------------------------

    document.add_heading(
        "Project Information",
        level=2,
    )

    document.add_paragraph(
        f"Name: {report.project.name}"
    )

    document.add_paragraph(
        f"Description: {report.project.description or '-'}"
    )

    document.add_paragraph(
        f"Created: {report.project.created_at}"
    )

    # -------------------------------------------------
    # Executive Summary
    # -------------------------------------------------

    document.add_heading(
        "Executive Summary",
        level=2,
    )

    document.add_paragraph(
        report.ai_summary.executive_summary
    )

    # -------------------------------------------------
    # Research Summary
    # -------------------------------------------------

    document.add_heading(
        "Research Summary",
        level=2,
    )

    document.add_paragraph(
        report.ai_summary.research_summary
    )

    # -------------------------------------------------
    # Productivity Insights
    # -------------------------------------------------

    document.add_heading(
        "Productivity Insights",
        level=2,
    )

    document.add_paragraph(
        report.ai_summary.productivity_insights
    )

    # -------------------------------------------------
    # Analytics
    # -------------------------------------------------

    analytics = report.analytics

    document.add_heading(
        "Analytics",
        level=2,
    )

    document.add_paragraph(
        f"Projects: {analytics.total_projects}"
    )

    document.add_paragraph(
        f"Documents: {analytics.total_documents}"
    )

    document.add_paragraph(
        f"Chat Sessions: {analytics.total_chat_sessions}"
    )

    document.add_paragraph(
        f"Messages: {analytics.total_messages}"
    )

    document.add_paragraph(
        f"Research Notes: {analytics.total_research_notes}"
    )

    # -------------------------------------------------
    # Documents
    # -------------------------------------------------

    document.add_heading(
        "Documents",
        level=2,
    )

    for item in report.documents:
        document.add_paragraph(
            item.filename,
            style="List Bullet",
        )

    # -------------------------------------------------
    # Chat Sessions
    # -------------------------------------------------

    document.add_heading(
        "Chat Sessions",
        level=2,
    )

    for session in report.chat_sessions:
        document.add_paragraph(
            f"{session.title} ({session.message_count} messages)",
            style="List Bullet",
        )

    # -------------------------------------------------
    # Research Notes
    # -------------------------------------------------

    document.add_heading(
        "Research Notes",
        level=2,
    )

    for note in report.research_notes:
        document.add_paragraph(
            note.title,
            style="List Bullet",
        )

    # -------------------------------------------------
    # Citations
    # -------------------------------------------------

    document.add_heading(
        "Citations",
        level=2,
    )

    for citation in report.citations:
        document.add_paragraph(
            f"{citation.filename} (Chunk {citation.chunk})",
            style="List Bullet",
        )

    buffer = BytesIO()

    document.save(buffer)

    return buffer.getvalue()